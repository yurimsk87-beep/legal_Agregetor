from __future__ import annotations

import argparse
import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup, Tag
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE_URL = "https://harant.ru"
LAWYERS_URL = f"{BASE_URL}/lawyers/"
WPDISCUZ_AJAX_URL = f"{BASE_URL}/wp-admin/admin-ajax.php"
SITEMAP_URL = f"{BASE_URL}/sitemap.xml"
SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_OUTPUT = SCRIPT_DIR / "outputs" / "harant_first_lawyer.xlsx"
DEFAULT_LIMIT = 1
MAX_REVIEW_PAGES = 20

SELECTORS = {
    "status": (".wrap_status.confirmed_ststus > span", ".confirmed_text_s"),
    "education": ".verification_education",
    "specialization": ".list_category",
    "prices": ".list_price li",
    "reviews": (".wpd-comment-text p",),
}
STATUS_NOISE_TEXT = ("Подтверждено Минюст РФ", "Юридический статус:")


class ParserError(RuntimeError):
    """Raised when Harant HTML no longer contains the expected data."""


@dataclass(frozen=True)
class LawyerRecord:
    name: str
    status: str
    description: str
    education: str
    specialization: str
    prices: str
    reviews: str
    profile_url: str
    source_url: str
    parsed_at: str


def build_session() -> requests.Session:
    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/125.0 Safari/537.36"
            ),
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "ru-RU,ru;q=0.9,en;q=0.8",
        }
    )
    return session


def fetch_html(session: requests.Session, url: str, timeout: int) -> BeautifulSoup:
    response = session.get(url, timeout=timeout)
    response.raise_for_status()
    return BeautifulSoup(response.text, "html.parser")


def fetch_text(session: requests.Session, url: str, timeout: int) -> str:
    response = session.get(url, timeout=timeout)
    response.raise_for_status()
    return response.text


def normalize_text(value: str) -> str:
    lines = []
    for raw_line in value.splitlines():
        line = re.sub(r"\s+", " ", raw_line).strip()
        if line:
            lines.append(line)
    return "\n".join(lines)


def element_text(element: Tag) -> str:
    return normalize_text(element.get_text("\n", strip=True))


def first_existing_text(soup: BeautifulSoup, selectors: Iterable[str]) -> str:
    for selector in selectors:
        element = soup.select_one(selector)
        if element:
            text = element_text(element)
            if text:
                return text
    return ""


def all_elements_text(soup: BeautifulSoup, selector: str) -> str:
    texts = [element_text(element) for element in soup.select(selector)]
    return "\n\n".join(text for text in texts if text)


def first_existing_elements_text(soup: BeautifulSoup, selectors: Iterable[str]) -> str:
    for selector in selectors:
        text = all_elements_text(soup, selector)
        if text:
            return text
    return ""


def extract_reviews_from_html(html: str) -> list[str]:
    soup = BeautifulSoup(html, "html.parser")
    return [element_text(element) for element in soup.select(SELECTORS["reviews"][0]) if element_text(element)]


def extract_wpdiscuz_post_id(profile_soup: BeautifulSoup) -> str:
    marker = '"wc_post_id":"'
    for script in profile_soup.find_all("script"):
        script_text = script.string or script.get_text()
        if marker in script_text:
            return script_text.split(marker, 1)[1].split('"', 1)[0]
    return ""


def extract_reviews(
    session: requests.Session,
    profile_soup: BeautifulSoup,
    profile_url: str,
    timeout: int,
) -> str:
    static_reviews = first_existing_elements_text(profile_soup, SELECTORS["reviews"])
    if static_reviews:
        return static_reviews

    post_id = extract_wpdiscuz_post_id(profile_soup)
    if not post_id:
        return ""

    reviews: list[str] = []
    seen_reviews: set[str] = set()
    offset = 0
    last_parent_id = "0"

    for _ in range(MAX_REVIEW_PAGES):
        response = session.post(
            WPDISCUZ_AJAX_URL,
            data={
                "action": "wpdLoadMoreComments",
                "postId": post_id,
                "offset": str(offset),
                "lastParentId": last_parent_id,
                "isFirstLoad": "1" if offset == 0 else "0",
                "wpdType": "",
            },
            headers={"Referer": profile_url},
            timeout=timeout,
        )
        response.raise_for_status()
        payload = response.json()
        data = payload.get("data", {}) if isinstance(payload, dict) else {}
        comment_list = data.get("comment_list", "")

        for review in extract_reviews_from_html(comment_list):
            if review not in seen_reviews:
                reviews.append(review)
                seen_reviews.add(review)

        if not data.get("is_show_load_more"):
            break

        next_last_parent_id = str(data.get("last_parent_id") or last_parent_id)
        if next_last_parent_id == last_parent_id and offset > 0:
            break

        last_parent_id = next_last_parent_id
        offset += 1

    return "\n\n".join(reviews)


def extract_education(profile_soup: BeautifulSoup) -> str:
    education_rows: list[str] = []
    for item in profile_soup.select(SELECTORS["education"]):
        year = element_text(item.select_one("span")) if item.select_one("span") else ""
        institution = element_text(item.select_one("p")) if item.select_one("p") else ""
        education = normalize_text(" ".join(part for part in (year, institution) if part))
        if education:
            education_rows.append(education)

    return "; ".join(education_rows)


def extract_prices(profile_soup: BeautifulSoup) -> str:
    price_rows: list[str] = []

    for item in profile_soup.select(SELECTORS["prices"]):
        spans = item.select("span")
        service_name = element_text(spans[0]) if len(spans) >= 1 else ""
        price = element_text(spans[1]) if len(spans) >= 2 else ""

        if not service_name:
            text_parts = [part.strip() for part in item.get_text("\n", strip=True).splitlines() if part.strip()]
            service_name = text_parts[0] if text_parts else ""
            price = text_parts[1] if len(text_parts) > 1 else price

        row = normalize_text(" - ".join(part for part in (service_name, price) if part))
        if row:
            price_rows.append(row)

    return "\n".join(price_rows)


def extract_status(profile_soup: BeautifulSoup) -> str:
    for selector in SELECTORS["status"]:
        raw_status = first_existing_text(profile_soup, (selector,))
        if not raw_status:
            continue

        status = raw_status
        for noise_text in STATUS_NOISE_TEXT:
            status = status.replace(noise_text, "")

        status = normalize_text(status)
        if status:
            return status

    return ""


def is_lawyer_profile_url(url: str) -> bool:
    parsed = urlparse(url)
    path_parts = [part for part in parsed.path.split("/") if part]

    if parsed.query or parsed.fragment:
        return False

    if len(path_parts) != 3:
        return False

    if path_parts[0] != "lawyers" or path_parts[1] == "page":
        return False

    return bool(path_parts[1] and path_parts[2])


def find_first_lawyer_profile_url(listing_soup: BeautifulSoup) -> str:
    profile_urls = find_lawyer_profile_urls(listing_soup, limit=1)
    if profile_urls:
        return profile_urls[0]

    raise ParserError("Не удалось найти ссылку на первую карточку юриста на /lawyers/.")


def find_lawyer_profile_urls(
    listing_soup: BeautifulSoup,
    limit: int,
    seen_urls: set[str] | None = None,
) -> list[str]:
    profile_urls: list[str] = []
    seen_urls = seen_urls if seen_urls is not None else set()

    for link in listing_soup.select(".name_card a[href]"):
        href = link.get("href", "")
        absolute_url = urljoin(BASE_URL, href)
        if not is_lawyer_profile_url(absolute_url) or absolute_url in seen_urls:
            continue
        profile_urls.append(absolute_url)
        seen_urls.add(absolute_url)
        if len(profile_urls) >= limit:
            return profile_urls

    for link in listing_soup.select('a[href*="/lawyers/"]'):
        href = link.get("href", "")
        absolute_url = urljoin(BASE_URL, href)
        if not is_lawyer_profile_url(absolute_url) or absolute_url in seen_urls:
            continue
        profile_urls.append(absolute_url)
        seen_urls.add(absolute_url)
        if len(profile_urls) >= limit:
            return profile_urls

    return profile_urls


def current_listing_page_number(url: str) -> int:
    match = re.search(r"/page/(\d+)/?", urlparse(url).path)
    return int(match.group(1)) if match else 1


def find_next_listing_url(listing_soup: BeautifulSoup, current_url: str) -> str:
    current_absolute_url = urljoin(BASE_URL, current_url)
    current_page = current_listing_page_number(current_absolute_url)

    for link in listing_soup.select("a.next.page-numbers[href], a[rel='next'][href]"):
        next_url = urljoin(BASE_URL, link.get("href", ""))
        if next_url and next_url != current_absolute_url:
            return next_url

    candidates: list[tuple[int, str]] = []
    for link in listing_soup.select('a[href*="/lawyers/page/"]'):
        next_url = urljoin(BASE_URL, link.get("href", ""))
        match = re.search(r"/lawyers/page/(\d+)/?", urlparse(next_url).path)
        if not match:
            continue
        page_number = int(match.group(1))
        if page_number > current_page:
            candidates.append((page_number, next_url))

    if candidates:
        return min(candidates)[1]

    return ""


def extract_load_more_payload(listing_soup: BeautifulSoup) -> dict[str, str]:
    button = listing_soup.select_one("#load_lawyers_search")
    if not button:
        return {}

    return {
        "action": "loadLawyersSearch",
        "max_num": button.get("data-max_num_s", "10"),
        "page_s": button.get("data-page_s", "2"),
        "city_s": button.get("data-city_s", ""),
        "ptype_s": button.get("data-ptype_s", ""),
        "get_s": button.get("data-get_s", ""),
        "market_ids": button.get("data-market_ids", ""),
        "onlin_list": button.get("data-onlin_list", ""),
    }


def load_more_lawyer_profile_urls(
    session: requests.Session,
    payload: dict[str, str],
    referer_url: str,
    limit: int,
    seen_urls: set[str],
    timeout: int,
) -> list[str]:
    profile_urls: list[str] = []

    while payload and len(profile_urls) < limit:
        response = session.post(
            WPDISCUZ_AJAX_URL,
            data=payload,
            headers={
                "Referer": referer_url,
                "X-Requested-With": "XMLHttpRequest",
            },
            timeout=timeout,
        )
        response.raise_for_status()
        data = response.json()
        if not data.get("success"):
            break

        html = data.get("data") or data.get("html") or ""
        ajax_soup = BeautifulSoup(html, "html.parser")
        remaining = limit - len(profile_urls)
        new_urls = find_lawyer_profile_urls(ajax_soup, remaining, seen_urls)
        if not new_urls:
            break

        profile_urls.extend(new_urls)

        next_page = str(data.get("paged") or "")
        if not next_page or next_page == "1":
            break

        payload["page_s"] = next_page
        payload["market_ids"] = str(data.get("market_ids") or "")
        if data.get("online_list"):
            payload["onlin_list"] = str(data.get("online_list") or "")

    return profile_urls


def extract_sitemap_locs(xml_text: str) -> list[str]:
    root = ET.fromstring(xml_text)
    return [element.text.strip() for element in root.findall(".//{*}loc") if element.text and element.text.strip()]


def collect_sitemap_lawyer_profile_urls(
    session: requests.Session,
    limit: int,
    seen_urls: set[str],
    timeout: int,
) -> list[str]:
    profile_urls: list[str] = []
    sitemap_locs = extract_sitemap_locs(fetch_text(session, SITEMAP_URL, timeout))
    property_sitemaps = [url for url in sitemap_locs if "property-sitemap" in url]

    for sitemap_url in property_sitemaps:
        if len(profile_urls) >= limit:
            break

        for profile_url in extract_sitemap_locs(fetch_text(session, sitemap_url, timeout)):
            if not is_lawyer_profile_url(profile_url) or profile_url in seen_urls:
                continue

            profile_urls.append(profile_url)
            seen_urls.add(profile_url)
            if len(profile_urls) >= limit:
                break

    return profile_urls


def collect_lawyer_profile_urls(
    session: requests.Session,
    lawyers_url: str,
    limit: int,
    timeout: int,
) -> list[str]:
    profile_urls: list[str] = []
    seen_urls: set[str] = set()
    seen_listing_urls: set[str] = set()
    listing_url = lawyers_url

    while listing_url and len(profile_urls) < limit:
        absolute_listing_url = urljoin(BASE_URL, listing_url)
        if absolute_listing_url in seen_listing_urls:
            break

        seen_listing_urls.add(absolute_listing_url)
        listing_soup = fetch_html(session, absolute_listing_url, timeout)
        remaining = limit - len(profile_urls)
        profile_urls.extend(find_lawyer_profile_urls(listing_soup, remaining, seen_urls))
        if len(profile_urls) >= limit:
            break

        load_more_payload = extract_load_more_payload(listing_soup)
        if load_more_payload:
            remaining = limit - len(profile_urls)
            profile_urls.extend(
                load_more_lawyer_profile_urls(
                    session,
                    load_more_payload,
                    absolute_listing_url,
                    remaining,
                    seen_urls,
                    timeout,
                )
            )
            if len(profile_urls) >= limit:
                break

        listing_url = find_next_listing_url(listing_soup, absolute_listing_url)

    if len(profile_urls) < limit:
        remaining = limit - len(profile_urls)
        profile_urls.extend(collect_sitemap_lawyer_profile_urls(session, remaining, seen_urls, timeout))

    return profile_urls


def parse_lawyer_profile(
    session: requests.Session,
    profile_soup: BeautifulSoup,
    profile_url: str,
    source_url: str,
    timeout: int,
) -> LawyerRecord:
    name = first_existing_text(
        profile_soup,
        (
            ".name_card [itemprop='name']",
            ".name_card h1",
            "h1[itemprop='name']",
            ".name_card",
            "h1",
        ),
    )
    if not name:
        raise ParserError("Не удалось найти имя юриста по селектору .name_card.")

    description = first_existing_text(
        profile_soup,
        (
            ".col_comment_description.is_open",
            ".col_comment_description",
            "[itemprop='description']",
        ),
    )
    if not description:
        raise ParserError(
            "Не удалось найти описание юриста по селектору "
            ".col_comment_description.is_open или .col_comment_description."
        )

    return LawyerRecord(
        name=name,
        status=extract_status(profile_soup),
        description=description,
        education=extract_education(profile_soup),
        specialization=all_elements_text(profile_soup, SELECTORS["specialization"]),
        prices=extract_prices(profile_soup),
        reviews=extract_reviews(session, profile_soup, profile_url, timeout),
        profile_url=profile_url,
        source_url=source_url,
        parsed_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    )


def save_to_excel(records: list[LawyerRecord], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Harant"

    headers = [
        "Имя",
        "Статус",
        "Описание",
        "Образование",
        "Специализация",
        "Услуги и стоимость",
        "Отзывы",
        "URL профиля",
        "Источник",
        "Дата парсинга",
    ]
    sheet.append(headers)
    for record in records:
        sheet.append(
            [
                record.name,
                record.status,
                record.description,
                record.education,
                record.specialization,
                record.prices,
                record.reviews,
                record.profile_url,
                record.source_url,
                record.parsed_at,
            ]
        )
    sheet.freeze_panes = "A2"

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = {
        "A": 34,
        "B": 18,
        "C": 80,
        "D": 60,
        "E": 48,
        "F": 48,
        "G": 70,
        "H": 58,
        "I": 32,
        "J": 22,
    }
    for column_idx, column_cells in enumerate(sheet.columns, start=1):
        letter = get_column_letter(column_idx)
        sheet.column_dimensions[letter].width = widths.get(letter, 20)
        for cell in column_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    workbook.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Парсер Harant: открывает /lawyers/, переходит в первую карточку "
            "юриста и сохраняет данные профиля в Excel."
        )
    )
    parser.add_argument(
        "--lawyers-url",
        default=LAWYERS_URL,
        help=f"Страница списка юристов. По умолчанию: {LAWYERS_URL}",
    )
    parser.add_argument(
        "--profile-url",
        default="",
        help="URL конкретного профиля юриста. Если указан, список /lawyers/ не используется.",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=DEFAULT_LIMIT,
        help=f"Сколько карточек юристов обработать. По умолчанию: {DEFAULT_LIMIT}",
    )
    parser.add_argument(
        "--output",
        default=str(DEFAULT_OUTPUT),
        help=f"Путь к Excel-файлу. По умолчанию: {DEFAULT_OUTPUT}",
    )
    parser.add_argument(
        "--timeout",
        type=int,
        default=30,
        help="Таймаут HTTP-запросов в секундах. По умолчанию: 30",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    output_path = Path(args.output).expanduser().resolve()
    limit = max(1, args.limit)

    session = build_session()
    if args.profile_url:
        profile_urls = [urljoin(BASE_URL, args.profile_url)]
        source_url = profile_urls[0]
    else:
        profile_urls = collect_lawyer_profile_urls(session, args.lawyers_url, limit, args.timeout)
        if not profile_urls:
            raise ParserError("Не удалось найти ссылки на карточки юристов на /lawyers/.")
        source_url = args.lawyers_url

    records: list[LawyerRecord] = []
    for index, profile_url in enumerate(profile_urls, start=1):
        print(f"[{index}/{len(profile_urls)}] {profile_url}")
        profile_soup = fetch_html(session, profile_url, args.timeout)
        records.append(
            parse_lawyer_profile(
                session,
                profile_soup,
                profile_url,
                source_url,
                args.timeout,
            )
        )

    save_to_excel(records, output_path)

    print(f"Готово: {output_path}")
    print(f"Обработано карточек: {len(records)}")
    if records:
        print(f"Первый юрист: {records[0].name}")
        print(f"Первый профиль: {records[0].profile_url}")


if __name__ == "__main__":
    main()
