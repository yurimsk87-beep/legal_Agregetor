from __future__ import annotations

import argparse
import re
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable
from urllib.parse import urlencode, urljoin, urlparse, urlunparse, parse_qsl

import requests
from bs4 import BeautifulSoup, Tag
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE_URL = "https://pravoved.ru"
API_BASE_URL = "https://api.pravoved.ru"
QUESTIONS_URL = f"{BASE_URL}/questions/"
SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_OUTPUT = SCRIPT_DIR / "outputs" / "pravoved_questions.xlsx"
DEFAULT_LIMIT = 15

LIST_SELECTORS = (
    "#app > div.Simple_root__QrLdg > main > div.Index_root__MfhUf > div > "
    "div.Index_main__JjK0r > ul.Questions_conteiner__NuVAh > li",
    "ul.Questions_conteiner__NuVAh > li",
    "ul[class*='Questions_conteiner'] > li",
)

CATEGORY_SELECTORS = (
    "div.QuestionListItem_labels__fc1IN > a > span",
    "div[class*='QuestionListItem_labels'] a > span",
    "div[class*='QuestionListItem_labels'] a",
)

TITLE_SELECTORS = (
    "a.QuestionListItem_title__pvdxT[href*='/question/']",
    "a[class*='QuestionListItem_title'][href*='/question/']",
    "article > div > a[href*='/question/']",
    "a[href*='/question/']",
)

QUESTION_TEXT_SELECTORS = (
    "div.QuestionListItem_content__AGUsU > div",
    "div[class*='QuestionListItem_content'] div[class*='QuestionListItem_text']",
    "div[class*='QuestionListItem_content'] div",
)

ANSWER_SELECTORS = (
    "#a6564290 > div > div.Answer_text__vhgJL",
    "[id^='a'] > div > div[class*='Answer_text']",
    "div[class*='Answer_text']",
)

DETAIL_TITLE_SELECTORS = (
    "#question-content h1",
    "h1[class*='Question_question__title']",
    "h1",
)

DETAIL_QUESTION_TEXT_SELECTORS = (
    "#question-content div[class*='Question_question__text']",
    "#question-content div[class*='Question_question__block']",
)

DETAIL_CATEGORY_SELECTORS = (
    "nav[aria-label='Хлебные крошки'] a[href^='/questions/']",
    "a[class*='Breadcrumbs_category__name'][href^='/questions/']",
)


class ParserError(RuntimeError):
    """Raised when Pravoved HTML no longer contains the expected data."""


@dataclass(frozen=True)
class QuestionRecord:
    category: str
    title: str
    question: str
    lawyer_answer: str
    question_url: str
    parsed_at: str


def build_session() -> requests.Session:
    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/125.0 Safari/537.36"
            ),
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "ru-RU,ru;q=0.9,en;q=0.8",
            "Cache-Control": "no-cache",
        }
    )
    return session


def fetch_html(session: requests.Session, url: str, timeout: int) -> BeautifulSoup:
    response = session.get(url, timeout=timeout)
    response.raise_for_status()
    return BeautifulSoup(response.text, "html.parser")


def fetch_json(session: requests.Session, url: str, referer: str, timeout: int, params: dict[str, str]) -> dict:
    response = session.get(
        url,
        headers={
            "Accept": "application/json, text/plain, */*",
            "Origin": BASE_URL,
            "Referer": referer,
        },
        params=params,
        timeout=timeout,
    )
    response.raise_for_status()
    payload = response.json()
    if not isinstance(payload, dict):
        raise ParserError(f"Pravoved API вернул неожиданный JSON для {url}.")
    return payload


def normalize_text(value: str) -> str:
    value = value.replace("\xa0", " ")
    lines = []
    for raw_line in value.splitlines():
        line = re.sub(r"\s+", " ", raw_line).strip()
        if line:
            lines.append(line)
    return "\n".join(lines)


def element_text(element: Tag | None) -> str:
    if not element:
        return ""
    return normalize_text(element.get_text("\n", strip=True))


def text_from_html_or_plain(value: object) -> str:
    text = "" if value is None else str(value)
    if "<" in text and ">" in text:
        return element_text(BeautifulSoup(text, "html.parser"))
    return normalize_text(text)


def first_existing_text(root: BeautifulSoup | Tag, selectors: Iterable[str]) -> str:
    for selector in selectors:
        element = root.select_one(selector)
        text = element_text(element)
        if text:
            return text
    return ""


def first_existing_elements(root: BeautifulSoup | Tag, selectors: Iterable[str]) -> list[Tag]:
    for selector in selectors:
        elements = root.select(selector)
        if elements:
            return elements
    return []


def is_question_detail_href(href: str) -> bool:
    return bool(re.search(r"/question/\d+/?$", href))


def question_listing_page_url(base_url: str, page: int) -> str:
    if page <= 1:
        return base_url

    parsed = urlparse(base_url)
    query = dict(parse_qsl(parsed.query, keep_blank_values=True))
    query["page"] = str(page)
    return urlunparse(parsed._replace(query=urlencode(query)))


def questions_url_categories(questions_url: str) -> list[str]:
    parsed = urlparse(questions_url)
    parts = [part for part in parsed.path.split("/") if part]
    if parts and parts[0] == "questions":
        parts = parts[1:]
    if parts and parts[-1].isdigit():
        parts = parts[:-1]
    return parts


def questions_url_params(questions_url: str) -> dict[str, str]:
    parsed = urlparse(questions_url)
    return {key: value for key, value in parse_qsl(parsed.query, keep_blank_values=True)}


def api_questions_page_url(categories: list[str], page: int) -> str:
    path_parts = [*categories]
    if page > 1:
        path_parts.append(str(page))
    path = "/".join(path_parts)
    return f"{API_BASE_URL}/api/questions/{path + '/' if path else ''}"


def find_question_link(article: Tag, page_url: str) -> tuple[str, str]:
    for selector in TITLE_SELECTORS:
        for link in article.select(selector):
            href = link.get("href") or ""
            title = element_text(link)
            if not title or re.fullmatch(r"\d+\s+\S*ответ\S*", title, re.IGNORECASE):
                continue
            if not is_question_detail_href(href):
                continue
            return title, urljoin(page_url, href)

    return "", ""


def extract_answer_texts(detail_soup: BeautifulSoup, max_answers: int) -> list[str]:
    answers: list[str] = []
    seen: set[str] = set()

    for selector in ANSWER_SELECTORS:
        for element in detail_soup.select(selector):
            text = element_text(element)
            if not text or text in seen:
                continue
            answers.append(text)
            seen.add(text)
            if max_answers > 0 and len(answers) >= max_answers:
                return answers
        if answers:
            return answers

    return answers


def fetch_lawyer_answer(
    session: requests.Session,
    question_url: str,
    source_url: str,
    timeout: int,
    max_answers: int,
) -> str:
    if not question_url:
        return ""

    try:
        response = session.get(question_url, headers={"Referer": source_url}, timeout=timeout)
        response.raise_for_status()
    except requests.RequestException as error:
        print(f"  ! Не удалось открыть страницу ответа: {question_url} ({error})")
        return ""

    detail_soup = BeautifulSoup(response.text, "html.parser")
    answers = extract_answer_texts(detail_soup, max_answers)
    return "\n\n".join(answers)


def has_lawyer_answer(record: QuestionRecord | None) -> bool:
    return bool(record and record.lawyer_answer.strip())


def log_skip_without_answer(record: QuestionRecord | None, question_url: str = "") -> None:
    url = record.question_url if record else question_url
    print(f"  - пропуск: нет ответа юриста{f' ({url})' if url else ''}")


def extract_detail_category(detail_soup: BeautifulSoup) -> str:
    for selector in DETAIL_CATEGORY_SELECTORS:
        links = detail_soup.select(selector)
        categories = [
            element_text(link)
            for link in links
            if link.get("href") not in ("/questions/", QUESTIONS_URL) and element_text(link)
        ]
        if categories:
            return categories[-1]
    return ""


def parse_question_detail(
    session: requests.Session,
    question_url: str,
    timeout: int,
    max_answers: int,
) -> QuestionRecord:
    absolute_url = urljoin(BASE_URL, question_url)
    response = session.get(absolute_url, headers={"Referer": QUESTIONS_URL}, timeout=timeout)
    response.raise_for_status()
    detail_soup = BeautifulSoup(response.text, "html.parser")

    title = first_existing_text(detail_soup, DETAIL_TITLE_SELECTORS)
    if not title:
        raise ParserError(f"Не удалось найти заголовок вопроса на странице {absolute_url}.")

    question_text = first_existing_text(detail_soup, DETAIL_QUESTION_TEXT_SELECTORS) or title
    answers = extract_answer_texts(detail_soup, max_answers)

    return QuestionRecord(
        category=extract_detail_category(detail_soup),
        title=title,
        question=question_text,
        lawyer_answer="\n\n".join(answers),
        question_url=absolute_url,
        parsed_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    )


def parse_question_item(
    session: requests.Session,
    item: Tag,
    page_url: str,
    timeout: int,
    max_answers: int,
    delay: float,
) -> QuestionRecord | None:
    article = item.select_one("article") or item
    title, question_url = find_question_link(article, page_url)
    if not title:
        return None

    category = first_existing_text(article, CATEGORY_SELECTORS)
    question_text = first_existing_text(article, QUESTION_TEXT_SELECTORS) or title

    if delay > 0:
        time.sleep(delay)

    lawyer_answer = fetch_lawyer_answer(session, question_url, page_url, timeout, max_answers)

    return QuestionRecord(
        category=category,
        title=title,
        question=question_text,
        lawyer_answer=lawyer_answer,
        question_url=question_url,
        parsed_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    )


def parse_question_api_item(
    session: requests.Session,
    item: dict,
    source_url: str,
    timeout: int,
    max_answers: int,
    delay: float,
) -> QuestionRecord | None:
    question_id = item.get("id")
    title = text_from_html_or_plain(item.get("title"))
    question_text = text_from_html_or_plain(item.get("text")) or title
    if not title and not question_text:
        return None

    raw_url = item.get("url") or (f"/question/{question_id}/" if question_id else "")
    question_url = urljoin(BASE_URL, str(raw_url))
    if not question_url:
        return None

    category = ""
    raw_category = item.get("category")
    if isinstance(raw_category, dict):
        category = text_from_html_or_plain(raw_category.get("name"))
        if category.lower() == "все":
            category = ""

    if not category:
        labels = item.get("labels")
        if isinstance(labels, list):
            for label in labels:
                if isinstance(label, dict):
                    category = text_from_html_or_plain(label.get("title") or label.get("name"))
                    if category:
                        break

    answer_count = int(item.get("answerCount") or 0)
    lawyer_answer = ""
    if answer_count > 0:
        if delay > 0:
            time.sleep(delay)
        lawyer_answer = fetch_lawyer_answer(session, question_url, source_url, timeout, max_answers)

    return QuestionRecord(
        category=category,
        title=title,
        question=question_text,
        lawyer_answer=lawyer_answer,
        question_url=question_url,
        parsed_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    )


def collect_api_question_records(
    session: requests.Session,
    questions_url: str,
    page: int,
    timeout: int,
    max_answers: int,
    delay: float,
    seen_urls: set[str],
    limit: int,
) -> list[QuestionRecord]:
    categories = questions_url_categories(questions_url)
    params = questions_url_params(questions_url)
    api_url = api_questions_page_url(categories, page)
    payload = fetch_json(session, api_url, questions_url, timeout, params)
    items = payload.get("items")
    if not isinstance(items, dict):
        return []

    records: list[QuestionRecord] = []
    for item in items.values():
        if len(records) >= limit:
            break
        if not isinstance(item, dict):
            continue

        raw_url = item.get("url") or (f"/question/{item.get('id')}/" if item.get("id") else "")
        question_url = urljoin(BASE_URL, str(raw_url))
        if not question_url or question_url in seen_urls:
            continue

        seen_urls.add(question_url)
        print(f"[{len(seen_urls)}] {question_url}")
        record = parse_question_api_item(session, item, questions_url, timeout, max_answers, delay)
        if has_lawyer_answer(record):
            records.append(record)
        else:
            log_skip_without_answer(record, question_url)

    return records


def collect_questions(
    session: requests.Session,
    questions_url: str,
    limit: int,
    pages: int,
    timeout: int,
    max_answers: int,
    delay: float,
) -> list[QuestionRecord]:
    records: list[QuestionRecord] = []
    seen_urls: set[str] = set()

    for page in range(1, max(1, pages) + 1):
        if len(records) >= limit:
            break

        page_url = question_listing_page_url(questions_url, page)
        if page == 1:
            print(f"[page {page}] {page_url}")
            listing_soup = fetch_html(session, page_url, timeout)
            items = first_existing_elements(listing_soup, LIST_SELECTORS)
            if not items:
                raise ParserError("Не удалось найти список вопросов Pravoved по селектору ul.Questions_conteiner__NuVAh > li.")

            for item in items:
                if len(records) >= limit:
                    break

                article = item.select_one("article") or item
                title, question_url = find_question_link(article, page_url)
                if not question_url or question_url in seen_urls:
                    continue

                seen_urls.add(question_url)
                print(f"[{len(records) + 1}/{limit}] {question_url}")
                record = parse_question_item(session, item, page_url, timeout, max_answers, delay)
                if has_lawyer_answer(record):
                    records.append(record)
                else:
                    log_skip_without_answer(record, question_url)
        else:
            print(f"[button page {page}] {api_questions_page_url(questions_url_categories(questions_url), page)}")
            remaining = limit - len(records)
            records.extend(
                collect_api_question_records(
                    session,
                    questions_url,
                    page,
                    timeout,
                    max_answers,
                    delay,
                    seen_urls,
                    remaining,
                )
            )

    return records


def save_to_excel(records: list[QuestionRecord], output_path: Path, include_meta: bool) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Pravoved"

    headers = ["Категория вопроса", "Заголовок вопроса", "Вопрос", "Ответ юриста"]
    if include_meta:
        headers.extend(["URL вопроса", "Дата парсинга"])

    sheet.append(headers)
    for record in records:
        row = [
            record.category,
            record.title,
            record.question,
            record.lawyer_answer,
        ]
        if include_meta:
            row.extend([record.question_url, record.parsed_at])
        sheet.append(row)

    sheet.freeze_panes = "A2"

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = {
        "A": 28,
        "B": 70,
        "C": 80,
        "D": 90,
        "E": 42,
        "F": 22,
    }
    for column_idx, column_cells in enumerate(sheet.columns, start=1):
        letter = get_column_letter(column_idx)
        sheet.column_dimensions[letter].width = widths.get(letter, 20)
        for cell in column_cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    workbook.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Парсер Pravoved: собирает вопросы с /questions/ и ответы юристов со страниц вопросов в Excel."
    )
    parser.add_argument(
        "--questions-url",
        default=QUESTIONS_URL,
        help=f"Страница списка вопросов. По умолчанию: {QUESTIONS_URL}",
    )
    parser.add_argument(
        "--question-url",
        action="append",
        default=[],
        help="URL конкретного вопроса. Можно указать несколько раз. Если указан, список /questions/ не используется.",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=DEFAULT_LIMIT,
        help=f"Сколько вопросов обработать. По умолчанию: {DEFAULT_LIMIT}",
    )
    parser.add_argument(
        "--pages",
        type=int,
        default=1,
        help="Сколько страниц списка вопросов просмотреть. По умолчанию: 1",
    )
    parser.add_argument(
        "--max-answers",
        type=int,
        default=1,
        help="Сколько ответов юристов брать с одной страницы вопроса. 0 = все найденные ответы. По умолчанию: 1",
    )
    parser.add_argument(
        "--output",
        default=str(DEFAULT_OUTPUT),
        help=f"Путь к Excel-файлу. По умолчанию: {DEFAULT_OUTPUT}",
    )
    parser.add_argument(
        "--include-meta",
        action="store_true",
        help="Добавить в Excel URL вопроса и дату парсинга.",
    )
    parser.add_argument(
        "--timeout",
        type=int,
        default=30,
        help="Таймаут HTTP-запросов в секундах. По умолчанию: 30",
    )
    parser.add_argument(
        "--delay",
        type=float,
        default=0.3,
        help="Пауза перед загрузкой страницы вопроса в секундах. По умолчанию: 0.3",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    limit = max(1, args.limit)
    output_path = Path(args.output).expanduser().resolve()

    session = build_session()
    if args.question_url:
        records = []
        for index, question_url in enumerate(args.question_url[:limit], start=1):
            print(f"[{index}/{min(limit, len(args.question_url))}] {question_url}")
            record = parse_question_detail(session, question_url, args.timeout, args.max_answers)
            if has_lawyer_answer(record):
                records.append(record)
            else:
                log_skip_without_answer(record, question_url)
            if args.delay > 0 and index < len(args.question_url[:limit]):
                time.sleep(args.delay)
    else:
        records = collect_questions(
            session=session,
            questions_url=args.questions_url,
            limit=limit,
            pages=args.pages,
            timeout=args.timeout,
            max_answers=args.max_answers,
            delay=max(0, args.delay),
        )

    if not records:
        raise ParserError("Не удалось собрать вопросы Pravoved с ответами юриста.")

    save_to_excel(records, output_path, include_meta=args.include_meta)
    with_answers = sum(1 for record in records if record.lawyer_answer)

    print(f"Готово: {output_path}")
    print(f"Обработано вопросов: {len(records)}")
    print(f"Вопросов с ответом юриста: {with_answers}")


if __name__ == "__main__":
    main()
