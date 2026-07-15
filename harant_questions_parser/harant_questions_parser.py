from __future__ import annotations

import argparse
import re
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable, Literal
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup, Tag
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE_URL = "https://harant.ru"
QUESTIONS_URL = f"{BASE_URL}/questions/"
AJAX_URL = f"{BASE_URL}/wp-admin/admin-ajax.php"
SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_OUTPUT = SCRIPT_DIR / "outputs" / "harant_questions.xlsx"
DEFAULT_LIMIT = 15

QUESTION_URL_PATTERN = re.compile(r"^/questions/q-\d+/?$")

CATEGORY_SELECTORS = (
    ".topic-title-wrap .cat_block a",
    ".cat_block a",
    "div.topic-title-wrap div.cat_block > a",
)

QUESTION_SELECTORS = (
    ".topic-content-wrap.p-rigth-30 > div > p",
    ".topic-content-wrap.p-rigth-30 p",
    ".topic-content-wrap > div > p",
    ".topic-content-wrap p",
    ".topic-content-wrap",
)

ANSWER_SELECTORS = (
    ".answer-user-data.p-rigth-30 .user-data",
    ".answer-user-data .user-data",
    "[class*='answer-excerpt']",
    ".answer-body .user-data",
    ".answer-body",
)


class ParserError(RuntimeError):
    """Raised when Harant HTML no longer contains the expected data."""


@dataclass(frozen=True)
class QuestionRecord:
    category: str
    question: str
    lawyer_answer: str
    question_url: str
    parsed_at: str


def build_session() -> requests.Session:
    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/125.0 Safari/537.36"
            ),
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "ru-RU,ru;q=0.9,en;q=0.8",
            "Cache-Control": "no-cache",
        }
    )
    return session


def normalize_text(value: str) -> str:
    value = value.replace("\xa0", " ")
    lines = []
    for raw_line in value.splitlines():
        line = re.sub(r"\s+", " ", raw_line).strip()
        if line:
            lines.append(line)
    return "\n".join(lines)


def element_text(element: Tag | None) -> str:
    if element is None:
        return ""
    return normalize_text(element.get_text("\n", strip=True))


def all_elements_text(root: BeautifulSoup | Tag, selector: str) -> str:
    texts = [element_text(element) for element in root.select(selector)]
    return "\n\n".join(text for text in texts if text)


def first_existing_text(root: BeautifulSoup | Tag, selectors: Iterable[str]) -> str:
    for selector in selectors:
        text = all_elements_text(root, selector)
        if text:
            return text
    return ""


def fetch_soup(session: requests.Session, url: str, timeout: int) -> BeautifulSoup:
    response = session.get(url, timeout=timeout)
    response.raise_for_status()
    return BeautifulSoup(response.text, "html.parser")


def is_question_detail_url(url: str) -> bool:
    parsed = urlparse(url)
    return bool(QUESTION_URL_PATTERN.match(parsed.path))


def normalize_question_url(page_url: str, href: str) -> str:
    return urljoin(page_url, href.split("#", 1)[0])


def find_question_urls(soup: BeautifulSoup | Tag, page_url: str) -> list[str]:
    urls: list[str] = []
    seen: set[str] = set()

    for link in soup.select("a[href]"):
        href = link.get("href", "")
        absolute_url = normalize_question_url(page_url, href)
        if not is_question_detail_url(absolute_url):
            continue
        if absolute_url in seen:
            continue
        seen.add(absolute_url)
        urls.append(absolute_url)

    return urls


def extract_load_more_payload(soup: BeautifulSoup) -> dict[str, str]:
    button = soup.select_one("#load_more_questions")
    if not button:
        return {}

    payload: dict[str, str] = {}
    for key, value in button.attrs.items():
        if not key.startswith("data-"):
            continue
        payload_key = key.removeprefix("data-").replace("-", "_")
        payload[payload_key] = str(value)

    payload.setdefault("action", "load_more_questions")
    return payload


def html_from_ajax_payload(payload: object) -> tuple[str, str | None]:
    if isinstance(payload, dict):
        html = payload.get("data") or payload.get("html") or payload.get("content") or ""
        current_page = payload.get("current_page")
        if isinstance(html, dict):
            html = html.get("html") or html.get("content") or ""
        return str(html), str(current_page) if current_page is not None else None

    return str(payload or ""), None


def load_more_questions(
    session: requests.Session,
    payload: dict[str, str],
    referer: str,
    timeout: int,
) -> tuple[BeautifulSoup | None, dict[str, str]]:
    response = session.post(
        AJAX_URL,
        data=payload,
        headers={
            "Accept": "application/json, text/javascript, */*; q=0.01",
            "Origin": BASE_URL,
            "Referer": referer,
            "X-Requested-With": "XMLHttpRequest",
        },
        timeout=timeout,
    )
    response.raise_for_status()

    if "application/json" in response.headers.get("content-type", ""):
        html, current_page = html_from_ajax_payload(response.json())
    else:
        html, current_page = response.text, None

    if current_page:
        payload["current_page"] = current_page
    elif payload.get("current_page", "").isdigit():
        payload["current_page"] = str(int(payload["current_page"]) + 1)

    html = html.strip()
    if not html or html == "0":
        return None, payload

    return BeautifulSoup(html, "html.parser"), payload


def collect_question_urls(
    session: requests.Session,
    start_url: str,
    limit: int,
    timeout: int,
    max_load_more: int,
) -> list[str]:
    soup = fetch_soup(session, start_url, timeout)
    urls: list[str] = []
    seen: set[str] = set()

    def add_urls(found_urls: Iterable[str]) -> int:
        added = 0
        for url in found_urls:
            if url in seen:
                continue
            seen.add(url)
            urls.append(url)
            added += 1
            if len(urls) >= limit:
                break
        return added

    add_urls(find_question_urls(soup, start_url))
    payload = extract_load_more_payload(soup)

    load_more_count = 0
    while len(urls) < limit and payload and load_more_count < max_load_more:
        if payload.get("current_page", "").isdigit() and payload.get("max_pages", "").isdigit():
            if int(payload["current_page"]) >= int(payload["max_pages"]):
                break

        loaded_soup, payload = load_more_questions(session, payload, start_url, timeout)
        load_more_count += 1
        if loaded_soup is None:
            break

        added = add_urls(find_question_urls(loaded_soup, start_url))
        if added == 0:
            break

    return urls[:limit]


def unique_texts(texts: Iterable[str]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for text in texts:
        normalized = normalize_text(text)
        if not normalized or normalized in seen:
            continue
        seen.add(normalized)
        result.append(normalized)
    return result


def extract_lawyer_answers(soup: BeautifulSoup, answer_mode: Literal["first", "all"]) -> str:
    answers: list[str] = []
    for answer_block in soup.select("[id^='answer-']"):
        answer_text = first_existing_text(answer_block, ANSWER_SELECTORS)
        if answer_text:
            answers.append(answer_text)

    answers = unique_texts(answers)
    if not answers:
        return ""
    if answer_mode == "first":
        return answers[0]
    return "\n\n---\n\n".join(answers)


def parse_question_detail(
    session: requests.Session,
    question_url: str,
    timeout: int,
    answer_mode: Literal["first", "all"],
) -> QuestionRecord:
    soup = fetch_soup(session, question_url, timeout)
    topic = soup.select_one("[id^='topic-']") or soup

    category = first_existing_text(topic, CATEGORY_SELECTORS)
    question = first_existing_text(topic, QUESTION_SELECTORS)
    lawyer_answer = extract_lawyer_answers(soup, answer_mode)

    if not question:
        raise ParserError(f"Не удалось найти текст вопроса на странице {question_url}.")

    return QuestionRecord(
        category=category,
        question=question,
        lawyer_answer=lawyer_answer,
        question_url=question_url,
        parsed_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    )


def save_to_excel(records: list[QuestionRecord], output_path: Path, include_meta: bool) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Harant Questions"

    headers = ["Категория вопроса", "Вопрос", "Ответ юриста"]
    if include_meta:
        headers.extend(["URL вопроса", "Дата парсинга"])

    sheet.append(headers)

    for record in records:
        row = [record.category, record.question, record.lawyer_answer]
        if include_meta:
            row.extend([record.question_url, record.parsed_at])
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = {
        "A": 28,
        "B": 80,
        "C": 95,
        "D": 45,
        "E": 22,
    }
    for column_letter, width in widths.items():
        sheet.column_dimensions[column_letter].width = width

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    workbook.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Парсер вопросов Harant: категория, вопрос и ответ юриста в Excel."
    )
    parser.add_argument("--limit", type=int, default=DEFAULT_LIMIT, help="Сколько карточек вопросов обработать.")
    parser.add_argument("--questions-url", default=QUESTIONS_URL, help="Страница списка вопросов Harant.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Путь к Excel-файлу результата.")
    parser.add_argument("--timeout", type=int, default=20, help="HTTP timeout в секундах.")
    parser.add_argument("--delay", type=float, default=0.4, help="Пауза между переходами в карточки.")
    parser.add_argument(
        "--max-load-more",
        type=int,
        default=100,
        help="Максимум AJAX-загрузок кнопки «Показать еще».",
    )
    parser.add_argument(
        "--answers",
        choices=("first", "all"),
        default="first",
        help="Сохранять первый ответ юриста или все уникальные ответы через разделитель.",
    )
    parser.add_argument(
        "--require-answer",
        action="store_true",
        help="Пропускать вопросы, где не найден ответ юриста.",
    )
    parser.add_argument(
        "--include-meta",
        action="store_true",
        help="Добавить в Excel URL вопроса и дату парсинга.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.limit < 1:
        raise SystemExit("--limit должен быть больше 0.")

    session = build_session()
    question_urls = collect_question_urls(
        session=session,
        start_url=args.questions_url,
        limit=args.limit,
        timeout=args.timeout,
        max_load_more=args.max_load_more,
    )

    if not question_urls:
        raise ParserError("Не удалось найти ссылки на карточки вопросов.")

    records: list[QuestionRecord] = []
    for index, question_url in enumerate(question_urls, start=1):
        print(f"[{index}/{len(question_urls)}] {question_url}")
        try:
            record = parse_question_detail(session, question_url, args.timeout, args.answers)
        except requests.RequestException as error:
            print(f"  HTTP ошибка: {error}")
            continue
        except ParserError as error:
            print(f"  Ошибка парсинга: {error}")
            continue

        if args.require_answer and not record.lawyer_answer:
            print("  Пропущено: ответ юриста не найден.")
        else:
            records.append(record)

        if args.delay:
            time.sleep(args.delay)

    if not records:
        raise ParserError("После обработки не осталось записей для Excel.")

    save_to_excel(records, args.output, args.include_meta)
    print(f"Готово: {args.output} ({len(records)} строк)")


if __name__ == "__main__":
    main()
