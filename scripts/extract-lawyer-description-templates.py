import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description="Extract lawyer description templates from an XLSX workbook.")
    parser.add_argument("input", help="Path to workbook with the 'Описание' column.")
    args = parser.parse_args()

    workbook_path = Path(args.input)
    if not workbook_path.exists():
        raise SystemExit(f"Workbook not found: {workbook_path}")

    workbook = load_workbook(workbook_path, data_only=True)
    worksheet = workbook.active
    headers = [normalize_header(worksheet.cell(1, column).value) for column in range(1, worksheet.max_column + 1)]

    try:
        description_column = headers.index("описание") + 1
    except ValueError as exc:
        raise SystemExit("Column not found: Описание") from exc

    templates = []
    for row_index in range(2, worksheet.max_row + 1):
        value = clean_text(worksheet.cell(row_index, description_column).value)
        if value:
            templates.append({"row": row_index, "description": value})

    print(json.dumps({"templates": templates}, ensure_ascii=False))


def normalize_header(value):
    return clean_text(value).casefold() if value else ""


def clean_text(value):
    if value is None:
        return ""

    text = str(value).replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ").strip()
    text = "\n".join(re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n"))
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text


if __name__ == "__main__":
    main()
