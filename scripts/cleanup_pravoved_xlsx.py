import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
XLSX_PATH = ROOT / "pravoved_parser" / "outputs" / "pravoved_1250_answered_questions_enriched_seo_fixed.xlsx"
JSON_PATH = ROOT / "pravoved_parser" / "outputs" / "pravoved_1250_answered_questions_site_import.json"


def normalized_key(value):
    text = "" if value is None else str(value)
    text = re.sub(r"\[контакт скрыт платформой\]", "", text.lower())
    text = re.sub(r"[^\wа-яё]+", " ", text, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", text).strip()


payload = json.loads(JSON_PATH.read_text(encoding="utf-8"))
rows = payload.get("rows") or []
row_by_question = {normalized_key(row.get("rawQuestion")): row for row in rows}

wb = load_workbook(XLSX_PATH)
ws = wb.active

rows_to_delete = []
seen_questions = set()

for row_index in range(2, ws.max_row + 1):
    question = ws.cell(row=row_index, column=1).value
    key = normalized_key(question)
    if not key:
        ws.cell(row=row_index, column=4).value = ""
        continue

    if key in seen_questions or key not in row_by_question:
        rows_to_delete.append(row_index)
        continue

    seen_questions.add(key)
    source_row = row_by_question[key]
    ws.cell(row=row_index, column=3).value = source_row.get("rawCategory", "")
    ws.cell(row=row_index, column=4).value = source_row.get("additionalCategory", "")
    ws.cell(row=row_index, column=5).value = source_row.get("rawTitle", "")
    ws.cell(row=row_index, column=6).value = source_row.get("lawyerAnswer", "")

for row_index in reversed(rows_to_delete):
    ws.delete_rows(row_index, 1)

wb.save(XLSX_PATH)

print(
    json.dumps(
        {
            "jsonRows": len(rows),
            "xlsxRows": ws.max_row - 1,
            "deletedRows": len(rows_to_delete),
            "nonemptyAdditional": sum(
                1
                for row_index in range(2, ws.max_row + 1)
                if str(ws.cell(row=row_index, column=4).value or "").strip()
            ),
        },
        ensure_ascii=False,
        indent=2,
    )
)
